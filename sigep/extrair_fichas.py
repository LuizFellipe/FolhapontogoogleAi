#!/usr/bin/env python3
"""
Script de Extração de Fichas Cadastrais (SIGEP / SEEDF) - Versão 2.0 Revisada
Lê PDFs no formato nome.matricula.pdf usando pdftotext -bbox-layout e BeautifulSoup.
Gera ficha.cadastral.DD.MM.YYYY.xlsx (multi-abas relacionais limpas) e ficha.cadastral.DD.MM.YYYY.json.
"""

import os
import re
import sys
import json
import subprocess
from datetime import datetime
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def get_pdf_bbox_layout(filepath):
    try:
        res = subprocess.run(
            ["pdftotext", "-bbox-layout", filepath, "-"],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            check=True
        )
        return res.stdout.decode("utf-8", errors="ignore")
    except Exception as e:
        print(f"Erro ao extrair bbox-layout de {filepath}: {e}", file=sys.stderr)
        return ""


def clean_val(val):
    if not val:
        return None
    val = val.strip()
    if val.lower() == "null" or val == "":
        return None
    return val


def parse_servidor(first_page, filename):
    servidor = {
        "arquivo_origem": filename,
        "matricula": None,
        "nome": None,
        "admissao": None,
        "cargo": None,
        "funcao": None,
        "ref_sal": None,
        "ch": None,
        "pcd": None,
        "reducao_ch": None,
        "readaptado": None,
        "identidade_funcional": None,
        "nascimento": None,
        "sexo": None,
        "cor_raca": None,
        "naturalidade": None,
        "nacionalidade": None,
        "uf_naturalidade": None,
        "ci_numero": None,
        "ci_orgao": None,
        "ci_uf": None,
        "ci_data_emissao": None,
        "cpf": None,
        "pis_pasep": None,
        "pis_emissao": None,
        "titulo_eleitoral": None,
        "titulo_zona": None,
        "titulo_secao": None,
        "estado_civil": None,
        "conjuge": None,
        "pai": None,
        "mae": None,
        "endereco": None,
        "bairro": None,
        "cidade": None,
        "uf_endereco": None,
        "cep": None,
        "telefones": [],
        "email": None,
        "especialidade_concurso": None,
        "escolaridade_salario": None
    }

    # Extrai matrícula do nome do arquivo como fallback de segurança
    mat_match = re.search(r'\.([0-9A-Za-z]+)\.pdf$', filename)
    if mat_match:
        servidor["matricula"] = mat_match.group(1)

    # Coleta todas as linhas e palavras com coordenadas
    lines_data = []
    for l in first_page.find_all("line"):
        words = [w.get_text() for w in l.find_all("word")]
        txt = " ".join(words).strip()
        if not txt:
            continue
        ly = float(l.get("ymin", 0))
        lx = float(l.get("xmin", 0))
        lines_data.append((ly, lx, txt))

    # pdftotext emite as linhas por bloco, não em ordem visual -> ordena por (y, x)
    lines_data.sort(key=lambda d: (d[0], d[1]))

    # Reconstrói linhas visuais: agrupa segmentos com y próximo (tolerância 3pt), ordenados por x
    rows = []
    for ly, lx, txt in lines_data:
        if rows and abs(rows[-1][0] - ly) <= 3:
            rows[-1][1].append((lx, txt))
        else:
            rows.append([ly, [(lx, txt)]])
    full_page_text = "\n".join(
        " ".join(t for _, t in sorted(segs)) for _, segs in rows
    )

    # 1. Nome, Matrícula, Admissão
    m_top = re.search(r'Nome:\s*(.*?)(?:\s+Matricula:\s*([0-9A-Za-z]+))(?:\s+Admissão:\s*([0-9/]+))?', full_page_text)
    if m_top:
        servidor["nome"] = clean_val(m_top.group(1))
        if clean_val(m_top.group(2)):
            servidor["matricula"] = clean_val(m_top.group(2))
        servidor["admissao"] = clean_val(m_top.group(3))

    # 2. Cargo, Função, Ref.Sal, CH
    m_cargo = re.search(r'Cargo:\s*(.*?)\s+Funçao:\s*(.*?)(?:\s*Ref\.Sal:\s*(.*?))?(?:\s+CH\s+([0-9]+h))?\s*$', full_page_text, re.MULTILINE)
    if m_cargo:
        servidor["cargo"] = clean_val(m_cargo.group(1))
        servidor["funcao"] = clean_val(m_cargo.group(2))
        servidor["ref_sal"] = clean_val(m_cargo.group(3))
        servidor["ch"] = clean_val(m_cargo.group(4))

    # 3. PCD, Redução CH
    m_pcd = re.search(r'Pessoa com Deficiência:\s*([^\n\r]+?)(?:\s+Redução de CH em\s+([^\n\r]+))?$', full_page_text, re.MULTILINE)
    if m_pcd:
        servidor["pcd"] = clean_val(m_pcd.group(1))
        if m_pcd.group(2):
            servidor["reducao_ch"] = clean_val(m_pcd.group(2))

    # 4. Readaptado, Identidade Funcional
    m_readapt = re.search(r'Readaptado:\s*([^\n\r]+?)(?:\s+Identidade Funcional\s+([^\n\r]+))?$', full_page_text, re.MULTILINE)
    if m_readapt:
        servidor["readaptado"] = clean_val(m_readapt.group(1))
        if m_readapt.group(2):
            servidor["identidade_funcional"] = clean_val(m_readapt.group(2))
    # Observação da readaptação ocupa várias linhas (y ~120..140) à direita do rótulo
    readapt_label = next((t for y, x, t in lines_data if t.startswith("Readaptado:")), None)
    if readapt_label:
        partes = [readapt_label.replace("Readaptado:", "").strip()]
        partes += [t for y, x, t in lines_data if 118 <= y < 140 and 90 <= x < 399]
        servidor["readaptado"] = clean_val(" ".join(p for p in partes if p))

    # 5. Nascimento, Sexo, Cor/Raça
    m_nasc = re.search(r'Nascimento:\s*([0-9/]+)\s+Sexo:\s*(.*?)\s*Cor/Raça:\s*([^\n\r]+?)(?:\s+Identidade Funcional\s+([^\n\r]+))?$', full_page_text, re.MULTILINE)
    if m_nasc:
        servidor["nascimento"] = clean_val(m_nasc.group(1))
        sexo_str = m_nasc.group(2)
        if "X MASCULINO" in sexo_str or "X  MASCULINO" in sexo_str:
            servidor["sexo"] = "MASCULINO"
        elif "X FEMININO" in sexo_str or "X  FEMININO" in sexo_str:
            servidor["sexo"] = "FEMININO"
        else:
            servidor["sexo"] = clean_val(sexo_str)
        servidor["cor_raca"] = clean_val(m_nasc.group(3))
        if m_nasc.group(4) and not servidor["identidade_funcional"]:
            servidor["identidade_funcional"] = clean_val(m_nasc.group(4))

    # 6. Naturalidade, Nacionalidade, UF
    m_nat = re.search(r'Naturalidade:\s*(.*?)\s*Nacionalidade:\s*(.*?)(?:\s+UF:\s*([A-Z]{2}))?$', full_page_text, re.MULTILINE)
    if m_nat:
        servidor["naturalidade"] = clean_val(m_nat.group(1))
        servidor["nacionalidade"] = clean_val(m_nat.group(2))
        servidor["uf_naturalidade"] = clean_val(m_nat.group(3))

    # 7. CI
    m_ci = re.search(r'CI/Número:\s*(.*?)\s*CI/Órgão:\s*(.*?)\s*CI/UF:\s*(.*?)\s*Data Emissão da CI:\s*([0-9/]*)', full_page_text)
    if m_ci:
        servidor["ci_numero"] = clean_val(m_ci.group(1))
        servidor["ci_orgao"] = clean_val(m_ci.group(2))
        servidor["ci_uf"] = clean_val(m_ci.group(3))
        servidor["ci_data_emissao"] = clean_val(m_ci.group(4))

    # 8. CPF, PIS/PASEP, Título Eleitoral
    m_doc = re.search(r'CPF:\s*([0-9]+)\s+PIS/PASEP:\s*([0-9]+)(?:\s+Emissão:\s*([0-9/]+))?\s+Título Eleitoral:\s*([0-9]+)(?:\s+Zona:\s*([0-9]+))?(?:\s+Seção:\s*([0-9]+))?', full_page_text)
    if m_doc:
        servidor["cpf"] = clean_val(m_doc.group(1))
        servidor["pis_pasep"] = clean_val(m_doc.group(2))
        servidor["pis_emissao"] = clean_val(m_doc.group(3))
        servidor["titulo_eleitoral"] = clean_val(m_doc.group(4))
        servidor["titulo_zona"] = clean_val(m_doc.group(5))
        servidor["titulo_secao"] = clean_val(m_doc.group(6))

    # 9. Estado Civil e Cônjuge (CORRIGIDO: isolamento por linha horizontal y ~ 200..210)
    ec_lines = [t for y, x, t in lines_data if 198 <= y <= 212]
    conjuge_val = None
    for item in ec_lines:
        if "Estado Civil:" in item:
            # ex: 'Estado Civil: CASADO Nome do Cônjuge: CONJUGE EXEMPLO' ou apenas 'Estado Civil: SOLTEIRO'
            m_ec = re.search(r'Estado Civil:\s*(.*?)(?:\s*Nome do Cônjuge:\s*(.*))?$', item)
            if m_ec:
                servidor["estado_civil"] = clean_val(m_ec.group(1))
                if m_ec.group(2):
                    conjuge_val = clean_val(m_ec.group(2))
        elif "Nome do Cônjuge:" in item:
            val_after = item.replace("Nome do Cônjuge:", "").strip()
            if val_after:
                conjuge_val = clean_val(val_after)
        elif not item.startswith("Estado Civil:") and not item.startswith("CPF:") and not item.startswith("PIS/PASEP:"):
            # Linha separada na mesma altura horizontal contendo o nome do cônjuge
            if not conjuge_val:
                conjuge_val = clean_val(item)

    # Proteção explícita contra invasão do Nome do Pai
    if conjuge_val and ("Nome do Pai:" in conjuge_val or "Filiação:" in conjuge_val):
        conjuge_val = None
    servidor["conjuge"] = conjuge_val

    # 10. Filiação: Pai e Mãe
    m_pai = re.search(r'Nome do Pai:\s*([^\n\r]*)', full_page_text)
    if m_pai:
        servidor["pai"] = clean_val(m_pai.group(1))
    m_mae = re.search(r'Nome da Mãe:\s*([^\n\r]*)', full_page_text)
    if m_mae:
        servidor["mae"] = clean_val(m_mae.group(1))

    # 11. Endereço
    m_end = re.search(r'Endereço:\s*(.*?)\s*Bairro:\s*(.*?)\s*Cidade:\s*(.*?)\s*UF:\s*([A-Z]{2})', full_page_text)
    if m_end:
        servidor["endereco"] = clean_val(m_end.group(1))
        servidor["bairro"] = clean_val(m_end.group(2))
        servidor["cidade"] = clean_val(m_end.group(3))
        servidor["uf_endereco"] = clean_val(m_end.group(4))

    m_cep = re.search(r'CEP:\s*([0-9]+)', full_page_text)
    if m_cep:
        servidor["cep"] = clean_val(m_cep.group(1))

    # 12. Telefones (faixa vertical y entre 255 e 275)
    for y, x, t in lines_data:
        if 255 <= y <= 275 and "Telefone:" in t:
            tels = re.findall(r'Telefone:\s*(\([0-9]{2}\)\s*[0-9]+|[0-9]{8,11})', t)
            for tel in tels:
                c_tel = clean_val(tel)
                if c_tel and c_tel not in servidor["telefones"]:
                    servidor["telefones"].append(c_tel)
        elif 255 <= y <= 275 and re.match(r'^\([0-9]{2}\)\s*[0-9]+$', t.strip()):
            c_tel = clean_val(t)
            if c_tel and c_tel not in servidor["telefones"]:
                servidor["telefones"].append(c_tel)

    m_email = re.search(r'Email:\s*([^\s\n\r]+@[^\s\n\r]+)', full_page_text)
    if m_email:
        servidor["email"] = clean_val(m_email.group(1))

    # Especialidade Concurso e Escolaridade
    m_esp = re.search(r'Especialidade/Disciplina de Concurso:[ \t]*([^\n\r]*)', full_page_text)
    if m_esp and m_esp.group(1).strip() not in ("/-", "/"):
        servidor["especialidade_concurso"] = clean_val(m_esp.group(1))
    m_esc = re.search(r'Escolaridade/Referência Salarial:[ \t]*([^\n\r]*)', full_page_text)
    if m_esc:
        servidor["escolaridade_salario"] = clean_val(m_esc.group(1))

    return servidor, lines_data


def split_by_column(page, x_split, y_min, y_max, x_min=0):
    """
    Divide cada <line> em segmentos esquerda/direita pelo xmin de cada palavra.
    Necessário porque o pdftotext às vezes funde as duas colunas numa única linha
    (ex: 'CRE/Unidade Administrativa: ... CRE/Unidade Administrativa: ...').
    Palavras com x < x_min são ignoradas.
    Retorna (esquerda, direita) como listas de texto ordenadas por (y, x).
    """
    left, right = [], []
    for l in page.find_all("line"):
        ly = float(l.get("ymin", 0))
        if not (y_min <= ly <= y_max):
            continue
        segs = {"L": [], "R": []}
        for w in l.find_all("word"):
            wx = float(w.get("xmin", 0))
            if wx < x_min:
                continue
            segs["L" if wx < x_split else "R"].append((wx, w.get_text()))
        for side, dest in (("L", left), ("R", right)):
            if segs[side]:
                txt = " ".join(t for _, t in segs[side]).strip()
                if txt:
                    dest.append((ly, segs[side][0][0], txt))
    left.sort(key=lambda d: (d[0], d[1]))
    right.sort(key=lambda d: (d[0], d[1]))
    return [t for _, _, t in left], [t for _, _, t in right]


def parse_cargas(page, lines_data, matricula):
    """
    Extrai Cargas Horárias separando rigorosamente Carga Principal (x < 290)
    e Carga Secundária (x >= 290) dentro da faixa 300 <= y <= 440.
    Elimina ruídos de cabeçalhos repetidos do JasperReports e descarta registros secundários vazios.
    """
    cargas = []
    # Seção de cargas termina no rótulo 'Especialidade/Disciplina de Concurso:'
    y_end = min((y for y, x, t in lines_data if t.startswith("Especialidade/")), default=441) - 1
    cp_lines, cs_lines = split_by_column(page, 290, 300, y_end)

    for tipo, lines in [("PRINCIPAL", cp_lines), ("SECUNDÁRIA", cs_lines)]:
        data = {
            "matricula": matricula,
            "tipo_carga": tipo,
            "unidade": "",
            "cre": "",
            "coord_externa": "",
            "lotacao": "",
            "turno": "",
            "atuacao": ""
        }
        current_key = None
        for line in lines:
            # Ignora linhas da seção inferior que entram na margem
            if (line.startswith("Especialidade/") or 
                line.startswith("Escolaridade/") or 
                re.match(r'^[0-9]{4}/', line) or 
                "POS-" in line or 
                "GRADUACAO" in line):
                continue

            m = re.match(r'^(Unidade:|CRE/Unidade Administrativa:|Coordenação Externa:|Lotação na Unidade:|Turno:|Atuação:)\s*(.*)$', line)
            if m:
                label = m.group(1)
                val = m.group(2).strip()
                if label == "Unidade:":
                    current_key = "unidade"
                    data["unidade"] = val
                elif label == "CRE/Unidade Administrativa:":
                    current_key = "cre"
                    # Remove duplicação do cabeçalho
                    val = re.sub(r'CRE/Unidade Administrativa:.*$', '', val).strip()
                    data["cre"] = val
                elif label == "Coordenação Externa:":
                    current_key = "coord_externa"
                    data["coord_externa"] = val
                elif label == "Lotação na Unidade:":
                    current_key = "lotacao"
                    data["lotacao"] = val
                elif label == "Turno:":
                    current_key = "turno"
                    data["turno"] = "" if val.lower() == "null" else val
                elif label == "Atuação:":
                    current_key = "atuacao"
                    data["atuacao"] = val
            else:
                if current_key and line.strip():
                    clean_line = re.sub(r'CRE/Unidade Administrativa:.*$', '', line).strip()
                    data[current_key] = (data[current_key] + " " + clean_line).strip()

        # Limpeza final de campos
        for k in data:
            if isinstance(data[k], str):
                v = data[k].strip()
                # Remove artefatos de quebra de coluna
                v = re.sub(r'\b(istrativa:|na:)\b', '', v).strip()
                if v.lower() == "null" or v == "":
                    data[k] = None
                else:
                    data[k] = v

        # Descarta linhas em branco ou cargas secundárias sem conteúdo real
        if data["unidade"] or data["atuacao"] or (data["turno"] and data["turno"] != "null"):
            cargas.append(data)

    return cargas


def parse_habilitacoes_e_componentes(page, matricula):
    """
    Extrai Habilitações (coluna esquerda x < 210, y: 448..565)
    e Componentes Curriculares (x >= 210, y: 448..565; sub-colunas em x=370) sem cortes.
    """
    habilitacoes = []
    componentes = []

    raw_hab, _ = split_by_column(page, 210, 448, 565)

    # Processa Habilitações
    # Lista separada por vírgula que quebra entre linhas (ex: 'GEOGRAFIA 1 E 2' + 'GRAU') -> une antes de dividir
    for item in " ".join(raw_hab).split(","):
        h = clean_val(item)
        if h and len(h) > 2:  # descarta fragmentos como 'EN', 'ED'
            habilitacoes.append({"matricula": matricula, "habilitacao": h})

    # Processa Componentes Curriculares: a área tem 2 sub-colunas
    # - esquerda (x < 370): um item por linha; nome de unidade longo quebra em 2 linhas
    #   (ex: 'CENTRO DE EDUC PROF ESCOLA TEC DO GUARA' + 'PROF TERESA ONDINA') -> une
    # - direita (x >= 370): lista separada por vírgula que quebra entre linhas -> une antes de dividir
    raw_lista, raw_virgula = split_by_column(page, 370, 448, 565, x_min=210)
    itens = []
    quebrado = False
    for line in raw_lista:
        if quebrado:
            itens[-1] += " " + line
            quebrado = False
        else:
            itens.append(line)
            quebrado = line.startswith("CENTRO DE")
    itens += " ".join(raw_virgula).split(",")
    vistos = set()
    for item in itens:
        c = clean_val(item)
        if c and c not in vistos:
            vistos.add(c)
            componentes.append({"matricula": matricula, "componente": c})

    return habilitacoes, componentes


def parse_cursos(soup, matricula):
    """
    Extrai Cursos e Certificados de Progressão de todas as páginas usando coordenadas.
    """
    from collections import defaultdict
    cursos = []
    seen = set()

    for page in soup.find_all("page"):
        row_groups = defaultdict(list)
        for l in page.find_all("line"):
            for w in l.find_all("word"):
                txt = w.get_text().strip()
                wx = float(w.get("xmin", 0))
                wy = float(w.get("ymin", 0))
                if 590 <= wy <= 795:
                    matched_k = None
                    for k in row_groups:
                        if abs(k - wy) < 4:
                            matched_k = k
                            break
                    if matched_k is None:
                        matched_k = wy
                    row_groups[matched_k].append((wx, txt))

        for k in sorted(row_groups.keys()):
            words_sorted = sorted(row_groups[k], key=lambda x: x[0])
            c_curso = " ".join(t for x, t in words_sorted if x < 220).strip()
            c_inst = " ".join(t for x, t in words_sorted if 220 <= x < 298).strip()
            c_emissao = " ".join(t for x, t in words_sorted if 298 <= x < 342).strip()
            c_util = " ".join(t for x, t in words_sorted if 342 <= x < 505).strip()
            c_data = " ".join(t for x, t in words_sorted if 505 <= x < 555).strip()
            c_ch = " ".join(t for x, t in words_sorted if x >= 555).strip()

            if c_emissao and c_data and c_ch.isdigit():
                curso_id = (matricula, c_curso, c_inst, c_emissao, c_util, c_data, int(c_ch))
                if curso_id not in seen:
                    seen.add(curso_id)
                    cursos.append({
                        "matricula": matricula,
                        "curso": c_curso,
                        "instituicao": c_inst,
                        "emissao": c_emissao,
                        "utilizacao": c_util,
                        "data_utilizacao": c_data,
                        "carga_horaria": int(c_ch)
                    })
    return cursos


def export_to_excel(servidores, cargas, habilitacoes, componentes, cursos, output_filename):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove default sheet

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    def style_sheet(ws):
        ws.views.sheetView[0].showGridLines = True
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row in range(2, ws.max_row + 1):
            for col in range(1, ws.max_column + 1):
                c = ws.cell(row=row, column=col)
                c.border = thin_border
                c.alignment = Alignment(vertical="center")

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(min(max_len + 3, 50), 12)

    # 1. SERVIDORES
    ws_serv = wb.create_sheet(title="Servidores")
    serv_cols = [
        ("Matrícula", "matricula"),
        ("Nome", "nome"),
        ("CPF", "cpf"),
        ("Admissão", "admissao"),
        ("Cargo", "cargo"),
        ("Função", "funcao"),
        ("Ref. Salarial", "ref_sal"),
        ("C.H.", "ch"),
        ("PCD", "pcd"),
        ("Redução CH", "reducao_ch"),
        ("Readaptado", "readaptado"),
        ("Identidade Funcional", "identidade_funcional"),
        ("Nascimento", "nascimento"),
        ("Sexo", "sexo"),
        ("Cor/Raça", "cor_raca"),
        ("Naturalidade", "naturalidade"),
        ("UF Naturalidade", "uf_naturalidade"),
        ("Nacionalidade", "nacionalidade"),
        ("CI Número", "ci_numero"),
        ("CI Órgão", "ci_orgao"),
        ("CI UF", "ci_uf"),
        ("CI Data Emissão", "ci_data_emissao"),
        ("PIS/PASEP", "pis_pasep"),
        ("PIS Emissão", "pis_emissao"),
        ("Título Eleitoral", "titulo_eleitoral"),
        ("Título Zona", "titulo_zona"),
        ("Título Seção", "titulo_secao"),
        ("Estado Civil", "estado_civil"),
        ("Cônjuge", "conjuge"),
        ("Nome do Pai", "pai"),
        ("Nome da Mãe", "mae"),
        ("Endereço", "endereco"),
        ("Bairro", "bairro"),
        ("Cidade", "cidade"),
        ("UF Endereço", "uf_endereco"),
        ("CEP", "cep"),
        ("Telefones", "telefones"),
        ("Email", "email"),
        ("Especialidade Concurso", "especialidade_concurso"),
        ("Escolaridade Salário", "escolaridade_salario"),
        ("Arquivo Origem", "arquivo_origem")
    ]
    ws_serv.append([c[0] for c in serv_cols])
    for s in servidores:
        row = []
        for _, key in serv_cols:
            val = s.get(key)
            if key == "telefones" and isinstance(val, list):
                row.append(" / ".join(val) if val else "")
            else:
                row.append(val or "")
        ws_serv.append(row)
    style_sheet(ws_serv)

    # 2. CARGAS HORÁRIAS
    ws_cargas = wb.create_sheet(title="Cargas_Horarias")
    cargas_cols = [
        ("Matrícula", "matricula"),
        ("Tipo Carga", "tipo_carga"),
        ("Unidade", "unidade"),
        ("CRE / Unidade Adm", "cre"),
        ("Coordenação Externa", "coord_externa"),
        ("Lotação na Unidade", "lotacao"),
        ("Turno", "turno"),
        ("Atuação", "atuacao")
    ]
    ws_cargas.append([c[0] for c in cargas_cols])
    for cg in cargas:
        ws_cargas.append([cg.get(k) or "" for _, k in cargas_cols])
    style_sheet(ws_cargas)

    # 3. CURSOS E PROGRESSÕES
    ws_cursos = wb.create_sheet(title="Cursos_Progressoes")
    cursos_cols = [
        ("Matrícula", "matricula"),
        ("Curso", "curso"),
        ("Instituição", "instituicao"),
        ("Emissão", "emissao"),
        ("Utilização", "utilizacao"),
        ("Data Utilização", "data_utilizacao"),
        ("Carga Horária (h)", "carga_horaria")
    ]
    ws_cursos.append([c[0] for c in cursos_cols])
    for cr in cursos:
        ws_cursos.append([cr.get(k) or "" for _, k in cursos_cols])
    style_sheet(ws_cursos)

    # 4. HABILITAÇÕES
    ws_hab = wb.create_sheet(title="Habilitacoes")
    ws_hab.append(["Matrícula", "Habilitação"])
    for h in habilitacoes:
        ws_hab.append([h["matricula"], h["habilitacao"]])
    style_sheet(ws_hab)

    # 5. COMPONENTES CURRICULARES
    ws_comp = wb.create_sheet(title="Componentes_Curriculares")
    ws_comp.append(["Matrícula", "Componente Curricular"])
    for comp in componentes:
        ws_comp.append([comp["matricula"], comp["componente"]])
    style_sheet(ws_comp)

    wb.save(output_filename)
    print(f"[OK] Planilha Excel salva com sucesso: {output_filename}")


def export_to_json(servidores, cargas, habilitacoes, componentes, cursos, output_filename):
    data = {
        "metadados": {
            "data_extracao": datetime.now().isoformat(),
            "total_servidores": len(servidores),
            "total_cargas": len(cargas),
            "total_cursos": len(cursos),
            "total_habilitacoes": len(habilitacoes),
            "total_componentes": len(componentes)
        },
        "tabelas": {
            "servidores": servidores,
            "cargas_horarias": cargas,
            "cursos_progressoes": cursos,
            "habilitacoes": habilitacoes,
            "componentes_curriculares": componentes
        }
    }
    with open(output_filename, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print(f"[OK] Dados JSON salvos com sucesso: {output_filename}")


def main():
    target_dir = sys.argv[1] if len(sys.argv) > 1 else "."
    all_files = os.listdir(target_dir)

    pdf_pattern = re.compile(r'^[a-zA-Z0-9_\-.]+\.([0-9A-Za-z]+)\.pdf$')
    pdf_files = [f for f in all_files if pdf_pattern.match(f) and not f.startswith("Emitir") and not f.startswith("Listagem")]
    pdf_files.sort()

    print(f"-> Localizados {len(pdf_files)} arquivos PDF de servidores para processamento.")

    all_servidores = []
    all_cargas = []
    all_habilitacoes = []
    all_componentes = []
    all_cursos = []

    for idx, fname in enumerate(pdf_files, 1):
        fpath = os.path.join(target_dir, fname)
        bbox_xml = get_pdf_bbox_layout(fpath)
        if not bbox_xml.strip():
            print(f"[{idx}/{len(pdf_files)}] [AVISO] Bbox vazio: {fname}")
            continue

        soup = BeautifulSoup(bbox_xml, "html.parser")
        first_page = soup.find("page")
        if not first_page:
            continue

        servidor, lines_data = parse_servidor(first_page, fname)
        cargas = parse_cargas(first_page, lines_data, servidor["matricula"])
        habs, comps = parse_habilitacoes_e_componentes(first_page, servidor["matricula"])
        cursos = parse_cursos(soup, servidor["matricula"])

        all_servidores.append(servidor)
        all_cargas.extend(cargas)
        all_habilitacoes.extend(habs)
        all_componentes.extend(comps)
        all_cursos.extend(cursos)

        if idx % 20 == 0 or idx == len(pdf_files):
            print(f"[{idx}/{len(pdf_files)}] Processados... ({servidor.get('nome') or fname})")

    today_str = datetime.now().strftime("%d.%m.%Y")
    excel_name = f"ficha.cadastral.{today_str}.xlsx"
    json_name = f"ficha.cadastral.{today_str}.json"

    export_to_excel(all_servidores, all_cargas, all_habilitacoes, all_componentes, all_cursos, os.path.join(target_dir, excel_name))
    export_to_json(all_servidores, all_cargas, all_habilitacoes, all_componentes, all_cursos, os.path.join(target_dir, json_name))

    print("\n--- RESUMO DA EXTRAÇÃO REVISADA ---")
    print(f"Servidores únicos extraídos: {len(all_servidores)}")
    print(f"Registros de Cargas Horárias: {len(all_cargas)}")
    print(f"Registros de Cursos/Progressões: {len(all_cursos)}")
    print(f"Registros de Habilitações: {len(all_habilitacoes)}")
    print(f"Registros de Componentes: {len(all_componentes)}")


if __name__ == "__main__":
    main()
